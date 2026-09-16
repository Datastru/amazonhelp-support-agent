"""
Converts a markdown report to .docx. Small hand-rolled converter, not a
full CommonMark implementation - just enough for the subset actually used
in report/REPORT.md and report/decision_log.md: headers, bold/italic/code
spans, bullet + numbered lists, tables, code fences, paragraphs.
"""
import argparse
import re

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_inline_runs(paragraph, text):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(chunk)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path, docx_path, title):
    doc = Document()
    doc.add_heading(title, level=0)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    n = len(lines)
    in_code = False
    code_lines = []

    while i < n:
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # headers
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        # tables: header row, separator row, then data rows
        if line.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|?[\s:-]+\|", lines[i + 1]):
            header = parse_table_row(line)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.paragraphs[0].add_run(text).bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    add_inline_runs(cell.paragraphs[0], text)
            continue

        def is_new_block(l):
            return (
                not l.strip()
                or re.match(r"^(#{1,3}\s|\s*-\s|\d+\.\s|\|)", l)
                or l.strip().startswith("```")
            )

        # bullet list (wrapped continuation lines are indented, no marker)
        m = re.match(r"^(\s*)-\s+(.*)", line)
        if m:
            indent = len(m.group(1))
            para_lines = [m.group(2)]
            i += 1
            while i < n and lines[i].strip() and not is_new_block(lines[i]):
                para_lines.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Bullet")
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.5)
            add_inline_runs(p, " ".join(para_lines))
            continue

        # numbered list
        m = re.match(r"^\d+\.\s+(.*)", line)
        if m:
            para_lines = [m.group(1)]
            i += 1
            while i < n and lines[i].strip() and not is_new_block(lines[i]):
                para_lines.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, " ".join(para_lines))
            continue

        # plain paragraph - accumulate wrapped lines until a blank line or a new block
        para_lines = [line]
        i += 1
        while i < n and lines[i].strip() and not is_new_block(lines[i]):
            para_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(l.strip() for l in para_lines))

    doc.save(docx_path)
    print(f"wrote {docx_path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md_path")
    ap.add_argument("docx_path")
    ap.add_argument("--title", default="Report")
    args = ap.parse_args()
    convert(args.md_path, args.docx_path, args.title)


if __name__ == "__main__":
    main()
